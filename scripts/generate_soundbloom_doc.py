# -*- coding: utf-8 -*-
"""Generate SoundBloom project overview as Word and PDF on Desktop."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

DESKTOP = Path.home() / "Desktop"
DOCX_PATH = DESKTOP / "SoundBloom_Описание_проекта.docx"
PDF_PATH = DESKTOP / "SoundBloom_Описание_проекта.pdf"
FONT_PATH = Path(r"C:\Windows\Fonts\arial.ttf")


def set_cell_shading(cell, fill_hex: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill_hex)
    cell._tc.get_or_add_tcPr().append(shading)


def add_table_docx(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_shading(hdr_cells[i], "1E3A5F")
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val
            for p in cells[ci].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()


def build_docx():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    title = doc.add_heading("SoundBloom — описание проекта", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        "Полный обзор музыкальной стриминговой платформы SoundBloom: "
        "польза, отличия от Spotify и других сервисов, особенности и статистика."
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Astana Polytechnic · Дипломный проект · 2026")
    doc.add_paragraph()

    doc.add_heading("Что такое SoundBloom", level=1)
    doc.add_paragraph(
        "SoundBloom — full-stack музыкальная стриминговая платформа (дипломный проект). "
        "Это не коммерческий конкурент Spotify, а полноценный рабочий продукт: "
        "веб-приложение с backend, базой данных, подписками, модерацией и студией для артистов."
    )

    doc.add_heading("В чём польза проекта", level=1)

    doc.add_heading("Для слушателя", level=2)
    for item in [
        "Один интерфейс для трендов, поиска, жанров, плейлистов и личной библиотеки.",
        "Гибридный каталог: превью из Jamendo и iTunes + собственные треки артистов SoundBloom.",
        "Библиотека в аккаунте: лайки, сохранённые альбомы, жанры, плейлисты.",
        "Тарифы Free / Premium / Family — без рекламы, lossless, офлайн-скачивание (Premium).",
        "3 языка: русский, казахский, английский.",
        "Тёмная и светлая тема, адаптивная вёрстка.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Для музыканта (Artist Studio)", level=2)
    for item in [
        "Регистрация артиста и загрузка своих треков (аудио, обложка, текст, теги, лицензия).",
        "Модерация перед публикацией (роли Operator / Admin).",
        "Аналитика: прослушивания, скачивания по каждому треку.",
        "Своя площадка для публикации музыки без посредников.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Для разработчика / защиты диплома", level=2)
    for item in [
        "71 API-эндпоинт, 24 таблицы БД, 10 backend-модулей.",
        "Полный цикл: auth → каталог → библиотека → подписки → промокоды → admin/operator → i18n.",
        "Стек: React 19 + Vite + TypeScript ↔ ASP.NET Core (C#) ↔ Supabase (PostgreSQL + Auth + Storage).",
        "Развёрнут и готов к демо (GitHub Pages + backend).",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Отличия от Spotify, Apple Music, Яндекс.Музыки", level=1)
    add_table_docx(
        doc,
        ["Аспект", "Spotify / Apple / Яндекс", "SoundBloom"],
        [
            ("Масштаб каталога", "Миллионы лицензированных треков", "Jamendo + iTunes превью + загрузки артистов"),
            ("Модель", "Закрытый коммерческий продукт", "Открытая full-stack платформа (свой код, своя БД)"),
            ("Для артистов", "Spotify for Artists (аналитика, без прямой загрузки)", "Artist Studio: загрузка → модерация → публикация"),
            ("Модерация", "Внутренняя, непрозрачная", "Явные роли Admin и Operator, approve/reject"),
            ("Локализация", "Много языков, kk часто слабо", "ru / kk / en из коробки + admin-панель переводов"),
            ("Подписки", "Реальные платежи, региональные цены", "Free / Premium / Family + промокоды"),
            ("Офлайн", "Полноценный офлайн на всех треках", "Скачивание только треков SoundBloom"),
            ("Реклама", "На free-тарифе", "На free: реклама + лимит 3 skip/день"),
            ("Lossless", "Spotify HiFi / Apple Lossless", "Заявлено на Premium"),
            ("Семейный план", "До 6 человек", "Family до 6 участников по username"),
        ],
        col_widths=[1.4, 2.2, 2.4],
    )

    doc.add_heading("Уникальные особенности SoundBloom", level=1)

    features = [
        (
            "1. Гибридный каталог «discovery + UGC»",
            "Три источника: Jamendo (независимая музыка), iTunes (превью популярных треков), "
            "SoundBloom (треки артистов через backend). Библиотека хранит лайки с указанием source "
            "(jamendo, itunes, soundbloom) — единая коллекция из разных источников.",
        ),
        (
            "2. Платформа «слушатель + артист» в одном продукте",
            "Spotify разделяет Spotify for Artists и Spotify Consumer. В SoundBloom один аккаунт "
            "может и слушать, и публиковать (после регистрации артиста).",
        ),
        (
            "3. Полный pipeline модерации",
            "Загрузка → moderation_requests (pending) → Operator/Admin approve/reject → трек в каталоге.",
        ),
        (
            "4. Локализация под Казахстан",
            "Казахский (kk) наравне с ru и en. Переводы управляются через admin API (/api/admin/translations).",
        ),
        (
            "5. Промокоды и семейный sharing",
            "Промокоды: percent / fixed / free_months, лимиты, сроки. "
            "Family: владелец добавляет участников по username (до 6).",
        ),
        (
            "6. Прозрачная архитектура",
            "Слои: Client (PWA) → ASP.NET (JWT, CORS) → Services → Supabase (RLS). "
            "Удобно показывать на защите и развивать дальше.",
        ),
    ]
    for title_text, body in features:
        p = doc.add_paragraph()
        run = p.add_run(title_text)
        run.bold = True
        doc.add_paragraph(body)

    doc.add_heading("Текст для защиты (краткий pitch)", level=1)
    doc.add_paragraph(
        "«SoundBloom решает проблему „шума“ в стриминге: реклама, сжатый звук, сложность для "
        "независимых артистов. Платформа объединяет discovery (Jamendo, iTunes), собственный "
        "UGC-каталог с модерацией, подписки без рекламы и трёхъязычный интерфейс. Цель диплома "
        "достигнута: система развёрнута, протестирована, содержит 71 REST-эндпоинт и полный "
        "пользовательский сценарий от регистрации до Premium.»"
    )

    doc.add_heading("Честные ограничения", level=1)
    for item in [
        "Не замена Spotify по объёму музыки — нет лицензий на мейнстрим-каталог.",
        "Jamendo/iTunes — в основном превью/стриминг через внешние API, не ваш контент.",
        "Офлайн и lossless — для своих треков и Premium, не для всего каталога.",
        "Платежи — демо-модель (сохранённая карта, промокоды), не полноценный payment gateway.",
        "Учебный/демо продукт, а не production-сервис с миллионами пользователей.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Статистика проекта", level=1)
    add_table_docx(
        doc,
        ["Метрика", "Значение"],
        [
            ("API endpoints", "71"),
            ("Таблицы БД", "24"),
            ("Языки UI", "3 (ru, kk, en)"),
            ("Роли", "user, artist, operator, admin"),
            ("Модули backend", "Auth, Catalog, Library, Artist, Subscription, Promo, Admin, Operator, I18n, Contact"),
            ("Frontend", "React 19, Vite, Tailwind, i18next, Framer Motion"),
            ("Backend", "ASP.NET Core, JWT, Supabase REST"),
            ("Деплой", "GitHub Pages (frontend) + backend API"),
        ],
        col_widths=[2.5, 3.5],
    )

    doc.add_heading("Архитектура (сквозной путь)", level=1)
    add_table_docx(
        doc,
        ["Слой", "Технологии", "Функции"],
        [
            ("Frontend", "React 19 · Vite · TypeScript · Tailwind", "Каталог, плеер, библиотека, artist studio, Premium, ru/kk/en"),
            ("Backend", "C# · ASP.NET Core · Supabase REST", "Каталог, стриминг, библиотека, подписки, промокоды, admin/operator"),
            ("База данных", "Supabase (PostgreSQL + Auth + Storage)", "Данные, авторизация, хранение файлов, RLS"),
        ],
        col_widths=[1.2, 2.0, 2.8],
    )

    doc.add_heading("Итог", level=1)
    doc.add_paragraph(
        "SoundBloom — не «ещё один Spotify», а full-stack музыкальная экосистема с акцентом на "
        "независимых артистов, локализацию ru/kk/en, модерацию контента и прозрачную архитектуру. "
        "Сильный дипломный проект с реальными product-фичами, но без масштаба и лицензий гигантов индустрии."
    )

    doc.save(DOCX_PATH)
    print(f"DOCX: {DOCX_PATH}")


def pdf_escape(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def build_pdf():
    pdfmetrics.registerFont(TTFont("Arial", str(FONT_PATH)))
    pdfmetrics.registerFont(TTFont("Arial-Bold", str(Path(r"C:\Windows\Fonts\arialbd.ttf"))))

    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Heading1"],
        fontName="Arial-Bold",
        fontSize=18,
        spaceAfter=12,
        alignment=1,
    )
    h1 = ParagraphStyle("H1", parent=styles["Heading2"], fontName="Arial-Bold", fontSize=14, spaceBefore=14, spaceAfter=8)
    h2 = ParagraphStyle("H2", parent=styles["Heading3"], fontName="Arial-Bold", fontSize=12, spaceBefore=10, spaceAfter=6)
    body = ParagraphStyle("Body", parent=styles["Normal"], fontName="Arial", fontSize=10, leading=14, spaceAfter=6)
    bullet = ParagraphStyle("Bullet", parent=body, leftIndent=14, bulletIndent=6, spaceAfter=4)

    story = []

    def add_h1(text):
        story.append(Paragraph(pdf_escape(text), h1))

    def add_h2(text):
        story.append(Paragraph(pdf_escape(text), h2))

    def add_p(text):
        story.append(Paragraph(pdf_escape(text), body))

    def add_bullet(text):
        story.append(Paragraph(f"• {pdf_escape(text)}", bullet))

    def add_table(headers, rows, col_widths):
        data = [headers] + rows
        t = Table(data, colWidths=col_widths, repeatRows=1)
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A5F")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Arial-Bold"),
                    ("FONTNAME", (0, 1), (-1, -1), "Arial"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        story.append(t)
        story.append(Spacer(1, 10))

    story.append(Paragraph(pdf_escape("SoundBloom — описание проекта"), title_style))
    add_p("Полный обзор музыкальной стриминговой платформы SoundBloom.")
    add_p("Astana Polytechnic · Дипломный проект · 2026")
    story.append(Spacer(1, 8))

    add_h1("Что такое SoundBloom")
    add_p(
        "SoundBloom — full-stack музыкальная стриминговая платформа (дипломный проект). "
        "Полноценный рабочий продукт с backend, БД, подписками, модерацией и студией для артистов."
    )

    add_h1("В чём польза проекта")
    add_h2("Для слушателя")
    for item in [
        "Один интерфейс для трендов, поиска, жанров, плейлистов и библиотеки.",
        "Каталог: Jamendo + iTunes + треки артистов SoundBloom.",
        "Библиотека: лайки, альбомы, жанры, плейлисты.",
        "Тарифы Free / Premium / Family.",
        "3 языка: ru, kk, en.",
    ]:
        add_bullet(item)

    add_h2("Для музыканта")
    for item in [
        "Загрузка треков с обложкой, текстом, тегами, лицензией.",
        "Модерация (Operator / Admin).",
        "Аналитика прослушиваний и скачиваний.",
    ]:
        add_bullet(item)

    add_h2("Для защиты диплома")
    for item in [
        "71 API, 24 таблицы, 10 backend-модулей.",
        "React + ASP.NET Core + Supabase.",
    ]:
        add_bullet(item)

    add_h1("Отличия от Spotify и других")
    add_table(
        ["Аспект", "Spotify / Apple / Яндекс", "SoundBloom"],
        [
            ["Масштаб", "Миллионы треков", "Jamendo + iTunes + UGC"],
            ["Модель", "Закрытый продукт", "Свой full-stack код"],
            ["Артисты", "Без прямой загрузки", "Artist Studio + модерация"],
            ["Локализация", "kk слабо", "ru / kk / en + admin"],
            ["Офлайн", "Весь каталог", "Только треки SoundBloom"],
            ["Free", "Реклама", "Реклама + 3 skip/день"],
        ],
        [3.2 * cm, 5.5 * cm, 6.3 * cm],
    )

    add_h1("Уникальные особенности")
    for title_text, text in [
        ("Гибридный каталог", "Jamendo + iTunes + собственные загрузки артистов."),
        ("Слушатель + артист", "Один аккаунт для прослушивания и публикации."),
        ("Модерация", "Pending → approve/reject → каталог."),
        ("Локализация KZ", "Казахский язык + admin API переводов."),
        ("Промокоды и Family", "Скидки и семейный план до 6 человек."),
    ]:
        add_h2(title_text)
        add_p(text)

    add_h1("Pitch для защиты")
    add_p(
        "SoundBloom решает проблему шума в стриминге и даёт артистам собственную площадку. "
        "71 REST-эндпоинт, модерация, подписки, ru/kk/en — система развёрнута и готова к демо."
    )

    add_h1("Ограничения")
    for item in [
        "Не замена Spotify по объёму каталога.",
        "Jamendo/iTunes — внешние API.",
        "Офлайн только для своих треков.",
        "Демо-платежи, не production.",
    ]:
        add_bullet(item)

    add_h1("Статистика")
    add_table(
        ["Метрика", "Значение"],
        [
            ["API endpoints", "71"],
            ["Таблицы БД", "24"],
            ["Языки UI", "ru, kk, en"],
            ["Роли", "user, artist, operator, admin"],
            ["Frontend", "React 19, Vite, Tailwind"],
            ["Backend", "ASP.NET Core, Supabase"],
        ],
        [5.5 * cm, 9.5 * cm],
    )

    add_h1("Итог")
    add_p(
        "SoundBloom — full-stack экосистема для слушателей и артистов с локализацией ru/kk/en "
        "и модерацией контента. Сильный дипломный проект с реальными product-фичами."
    )

    doc.build(story)
    print(f"PDF: {PDF_PATH}")


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print("Done.")
